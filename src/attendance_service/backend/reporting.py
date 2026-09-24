"""Optional Excel reporting sourced exclusively from PostgreSQL."""

from __future__ import annotations

from collections import defaultdict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from attendance_service.database.repository import AttendanceRepository

HEADER_FILL = PatternFill(
    start_color="1F4E78", end_color="1F4E78", fill_type="solid"
)
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
NORMAL_FONT = Font(name="Arial")
TITLE_FONT = Font(name="Arial", bold=True, size=14)


def _header(worksheet: Any, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = worksheet.cell(row=row, column=column, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def generate_excel_report(
    repository: AttendanceRepository, source_key: str, output_path: Path
) -> None:
    events = repository.fetch_report_events(source_key)
    invalid = repository.fetch_invalid_report_events(source_key)
    if not events and not invalid:
        raise ValueError(f"No database records found for source: {source_key}")

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview["A1"] = "Attendance Report - PostgreSQL"
    overview["A1"].font = TITLE_FONT
    overview.merge_cells("A1:B1")
    overview.append([])
    # Excel stores naive datetimes. The label records that this naive value is UTC.
    overview.append(
        ["Generated at (UTC):", datetime.now(UTC).replace(tzinfo=None)]
    )
    overview.append(["Source:", source_key])
    overview.append(["Valid attendance records:", len(events)])
    overview.append(["Rejected attendance records:", len(invalid)])
    overview.append(
        ["Unique employees:", len({str(row[0]) for row in events})]
    )
    if events:
        overview.append(["Oldest device timestamp:", events[0][2]])
        overview.append(["Newest device timestamp:", events[-1][2]])
    overview.column_dimensions["A"].width = 30
    overview.column_dimensions["B"].width = 26

    monthly = workbook.create_sheet("Monthly Summary")
    monthly["A1"] = "Monthly Summary"
    monthly["A1"].font = TITLE_FONT
    _header(
        monthly,
        3,
        ["Month", "Total Punches", "Unique Employees", "First Punch", "Last Punch"],
    )
    grouped: dict[str, list[tuple[Any, ...]]] = defaultdict(list)
    for row in events:
        grouped[row[2].strftime("%Y-%m")].append(row)
    for row_number, month in enumerate(sorted(grouped), start=4):
        values = grouped[month]
        monthly.append(
            [
                month,
                len(values),
                len({str(value[0]) for value in values}),
                values[0][2],
                values[-1][2],
            ]
        )
        monthly.cell(row_number, 4).number_format = "yyyy-mm-dd hh:mm:ss"
        monthly.cell(row_number, 5).number_format = "yyyy-mm-dd hh:mm:ss"
    for column, width in {"A": 12, "B": 16, "C": 18, "D": 20, "E": 20}.items():
        monthly.column_dimensions[column].width = width
    monthly.freeze_panes = "A4"

    raw = workbook.create_sheet("Raw Data")
    _header(
        raw,
        1,
        ["Employee ID", "Employee Name", "Punch Timestamp", "Status", "Punch Code"],
    )
    for row in events:
        raw.append(list(row))
        raw.cell(raw.max_row, 3).number_format = "yyyy-mm-dd hh:mm:ss"
    for column, width in {"A": 14, "B": 28, "C": 20, "D": 12, "E": 12}.items():
        raw.column_dimensions[column].width = width
    raw.freeze_panes = "A2"

    rejected = workbook.create_sheet("Rejected")
    _header(
        rejected,
        1,
        [
            "Employee ID",
            "Employee Name",
            "Punch Timestamp",
            "Status",
            "Punch Code",
            "Validation Error",
        ],
    )
    for row in invalid:
        rejected.append(list(row))
        rejected.cell(rejected.max_row, 3).number_format = "yyyy-mm-dd hh:mm:ss"
    rejected.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
