from datetime import datetime

from openpyxl import load_workbook

from attendance_service.backend.reporting import generate_excel_report


class Repository:
    def fetch_report_events(self, _):
        return [
            ("0007", "Employee", datetime(2026, 9, 12, 9, 15), "1", "0")
        ]

    def fetch_invalid_report_events(self, _):
        return [
            (
                "0007",
                "Employee",
                datetime(2119, 3, 17, 21, 31),
                "1",
                "0",
                "future_timestamp",
            )
        ]


def test_report_is_generated_from_repository_data(tmp_path) -> None:
    output = tmp_path / "report.xlsx"

    generate_excel_report(Repository(), "test-source", output)

    workbook = load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "Overview",
        "Monthly Summary",
        "Raw Data",
        "Rejected",
    ]
    assert workbook["Raw Data"]["A2"].value == "0007"
    assert workbook["Rejected"]["C2"].value.year == 2119

